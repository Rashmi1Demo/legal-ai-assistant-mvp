from io import BytesIO
from pathlib import Path
from typing import Iterable

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


class DocumentProcessingError(ValueError):
    pass


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r", "\n").split("\n")]
    cleaned = "\n".join(line for line in lines if line)
    return cleaned.strip()


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentProcessingError(
            f"Unsupported file type '{suffix or 'unknown'}'. Use PDF, DOCX, or TXT."
        )

    if not content:
        raise DocumentProcessingError(f"'{filename}' is empty.")

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise DocumentProcessingError(f"'{filename}' exceeds the 10 MB demo limit.")

    try:
        if suffix == ".txt":
            text = content.decode("utf-8-sig", errors="replace")

        elif suffix == ".docx":
            document = Document(BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)

        else:  # PDF
            reader = PdfReader(BytesIO(content))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)

    except Exception as exc:
        raise DocumentProcessingError(f"Could not extract text from '{filename}': {exc}") from exc

    text = _normalize_text(text)
    if not text:
        raise DocumentProcessingError(
            f"No readable text was extracted from '{filename}'. Scanned/image-only PDFs are not supported in this MVP."
        )

    return text


def combine_documents(documents: Iterable[tuple[str, str]]) -> str:
    sections = []
    for filename, text in documents:
        sections.append(f"===== DOCUMENT: {filename} =====\n{text}")
    return "\n\n".join(sections).strip()
